"""Questionnaire profiles without saving personal responses on disk.

The synthetic workbook supplies test inputs, never eligibility labels.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import math
from pathlib import Path
import re
from typing import Any, Iterator

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
EXCEL_SHEET = "Reponses_synthetiques"
EXCEL_HEADERS = ["User", *[f"Q{i}" for i in range(1, 135)]]
DATE_IDS = {"Q26", "Q27"}
INTEGER_IDS = {"Q2", "Q51", "Q52", "Q53"}
UNKNOWN = {"je ne sais pas", "à déterminer", "pas encore décidée", "non disponible"}


def load_schema() -> dict:
    """Load the questionnaire shipped with this project, independently of cwd."""
    with (DATA / "questionnaire_schema.json").open(encoding="utf-8") as handle:
        return json.load(handle)


def _missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: Any) -> int | float:
    if isinstance(value, bool):
        raise ValueError("une réponse Oui/Non n’est pas un montant")
    if isinstance(value, str):
        cleaned = re.sub(r"\s+", "", value).replace("\u00a0", "").replace("\u202f", "")
        # Accept French decimal commas; mixed comma/dot notation is ambiguous.
        if "," in cleaned and "." in cleaned:
            raise ValueError("utiliser une virgule décimale et des espaces pour les milliers")
        cleaned = cleaned.replace(",", ".")
    elif isinstance(value, (int, float, Decimal)):
        cleaned = str(value)
    else:
        raise ValueError("nombre attendu")
    try:
        result = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError("nombre illisible") from exc
    if not result.is_finite() or result < 0:
        raise ValueError("le nombre doit être fini et positif ou nul")
    if result == result.to_integral_value():
        return int(result)
    floating = float(result)
    if not math.isfinite(floating):
        raise ValueError("nombre trop grand")
    return floating


def _date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str):
        raise ValueError("date attendue au format AAAA-MM-JJ ou JJ/MM/AAAA")
    cleaned = value.strip()
    try:
        return datetime.fromisoformat(cleaned).date().isoformat()
    except ValueError:
        try:
            return datetime.strptime(cleaned, "%d/%m/%Y").date().isoformat()
        except ValueError as exc:
            raise ValueError("date invalide, utiliser AAAA-MM-JJ ou JJ/MM/AAAA") from exc


def normalize_answers(raw: dict, schema: dict | None = None) -> tuple[dict, list[str]]:
    """Validate Q IDs and types; keep missing responses unknown, not zero/Non.

    Unknown choice strings are retained with an explicit warning so imports do
    not silently discard user content. Invalid numbers/dates become None.
    """
    if not isinstance(raw, dict):
        raise TypeError("Les réponses doivent être un dictionnaire Q1…Q134.")
    schema = load_schema() if schema is None else schema
    definitions = {question["id"]: question for question in schema["questions"]}
    answers: dict[str, Any] = {}
    warnings: list[str] = []
    for key, value in raw.items():
        if key not in definitions:
            if key != "User":
                warnings.append(f"Champ ignoré : {key} ne fait pas partie du questionnaire.")
            continue
        question = definitions[key]
        if _missing(value):
            answers[key] = None
            continue
        try:
            if key in DATE_IDS:
                if isinstance(value, str) and value.strip().casefold() in UNKNOWN | {"non concerné"}:
                    answers[key] = value.strip()
                else:
                    answers[key] = _date(value)
            elif question["type"] == "number":
                number = _number(value)
                if key in INTEGER_IDS and not isinstance(number, int):
                    raise ValueError("un nombre entier est attendu")
                answers[key] = number
            elif question["type"] == "multi":
                if isinstance(value, str):
                    parts = value.split(";")
                elif isinstance(value, (list, tuple)) and all(isinstance(item, str) for item in value):
                    parts = list(value)
                else:
                    raise ValueError("liste de choix ou texte séparé par des points-virgules attendu")
                answers[key] = list(dict.fromkeys(part.strip() for part in parts if part.strip())) or None
                for item in answers[key] or []:
                    if item not in question.get("options", []):
                        warnings.append(f"{key} : choix non prévu conservé : {item}.")
            else:
                if not isinstance(value, (str, int, float, Decimal, date)) or isinstance(value, bool):
                    raise ValueError("texte attendu")
                if isinstance(value, (float, Decimal)) and not math.isfinite(float(value)):
                    raise ValueError("valeur non finie")
                text = value.isoformat() if isinstance(value, date) else str(value).strip()
                answers[key] = text
                if question["type"] == "single" and text not in question.get("options", []):
                    warnings.append(f"{key} : choix non prévu conservé : {text}.")
        except (ValueError, OverflowError) as exc:
            answers[key] = None
            warnings.append(f"{key} : {exc}. La réponse reste inconnue.")
    return answers, warnings


def _display(value: Any) -> str:
    if isinstance(value, list):
        return " ; ".join(str(item) for item in value)
    return str(value)


def _known(value: Any) -> bool:
    return not _missing(value) and value != [] and str(value).strip().casefold() not in UNKNOWN


def _logical_warnings(answers: dict) -> list[str]:
    warnings = []
    def greater(left: str, right: str) -> bool:
        a, b = answers.get(left), answers.get(right)
        return isinstance(a, (float, int)) and isinstance(b, (float, int)) and a > b
    if greater("Q32", "Q31"):
        warnings.append("Le montant recherché (Q32) dépasse le coût total du projet (Q31).")
    if greater("Q53", "Q52"):
        warnings.append("Les emplois stables (Q53) dépassent le total des emplois à créer (Q52).")
    if answers.get("Q9") == "Auto-entrepreneur" and answers.get("Q6") in {"Avec 1 associé", "Avec 2 associés", "Avec 3 associés ou plus"}:
        warnings.append("La forme Auto-entrepreneur (Q9) et les associés déclarés (Q6) nécessitent une clarification.")
    try:
        if date.fromisoformat(answers.get("Q27", "")) < date.fromisoformat(answers.get("Q26", "")):
            warnings.append("Le début d’activité (Q27) précède la création juridique (Q26) : vérifier les dates.")
    except (ValueError, TypeError):
        pass
    if answers.get("Q113") == "Oui" and answers.get("Q116") == "Oui":
        warnings.append("L’entrée d’un investisseur (Q113) est acceptée, mais l’ouverture du capital est refusée (Q116).")
    return warnings


def build_profile(answers: dict, schema: dict | None = None) -> dict:
    """Make labelled facts and consistency warnings, with no eligibility verdict."""
    schema = load_schema() if schema is None else schema
    normalized, warnings = normalize_answers(answers, schema)
    warnings.extend(_logical_warnings(normalized))
    definitions = {question["id"]: question for question in schema["questions"]}
    labelled = [{"id": q["id"], "label": q["label"], "value": normalized[q["id"]]}
                for q in schema["questions"] if _known(normalized.get(q["id"]))]
    missing = [f"{key} — {definitions[key]['label']}" for key in schema.get("essential_ids", [])
               if key in definitions and not _known(normalized.get(key))]
    summary = "\n".join(f"{item['id']} — {item['label']} : {_display(item['value'])}" for item in labelled)
    return {"answers": normalized, "labeled_answers": labelled,
            "missing_information": missing, "warnings": warnings, "summary": summary}


def build_query(profile: dict, schema: dict | None = None) -> str:
    """Retrieve using concise, positive project facts rather than every question."""
    answers = profile.get("answers", profile)
    fields = [("Q17", "Activité"), ("Q18", "Produit/service"), ("Q19", "Secteur"),
              ("Q25", "Stade"), ("Q33", "Objectif du financement"), ("Q9", "Forme juridique"),
              ("Q11", "Région"), ("Q31", "Coût du projet DH"), ("Q32", "Besoin DH"),
              ("Q40", "Apport personnel DH"), ("Q48", "Chiffre d’affaires DH"),
              ("Q52", "Emplois à créer"), ("Q115", "Financements acceptés")]
    pieces = [f"{label}: {_display(answers[key])}" for key, label in fields
              if _known(answers.get(key)) and str(answers[key]).casefold() not in {"non", "non concerné"}]
    positive_tags = {
        "Q4": "Marocain résidant à l’étranger MRE",
        "Q7": "entreprise dirigée par une femme",
        "Q14": "milieu rural",
        "Q71": "financement marché contrat obtenu",
        "Q78": "innovation nouveau produit service procédé",
        "Q79": "prototype existant",
        "Q80": "tests techniques avant commercialisation",
        "Q81": "recherche et développement R&D",
        "Q83": "création extension unité de production",
        "Q84": "industrialisation lancement pilote",
        "Q87": "panneaux solaires énergie renouvelable",
        "Q88": "efficacité énergétique",
        "Q89": "économie réutilisation de l’eau",
        "Q90": "recyclage valorisation déchets",
        "Q91": "réduction pollution émissions",
        "Q92": "fabrication équipements économie verte",
        "Q100": "entreprise exportatrice",
        "Q107": "prospection certification marketing international export",
        "Q117": "financement participatif sans intérêt classique",
        "Q122": "renforcement fonds propres financement long terme",
    }
    pieces.extend(tag for key, tag in positive_tags.items() if answers.get(key) == "Oui")
    if str(answers.get("Q5", "")).startswith("Retour définitif"):
        pieces.append(f"Ancien MRE : {answers['Q5']}")
    if answers.get("Q8") in {"Plus de 50 %", "100 %"}:
        pieces.append("capital détenu majoritairement par des femmes")
    if _known(answers.get("Q94")) and answers["Q94"] != "Non concerné":
        pieces.append(f"Activité touristique : {answers['Q94']}")
    if _known(answers.get("Q22")) and answers["Q22"] != "Aucun de ces secteurs":
        pieces.append(f"Secteur particulier déclaré : {answers['Q22']}")
    return "\n".join(pieces) or "Programmes de financement et accompagnement des entreprises au Maroc"


def iter_excel_profiles(source: Path | BytesIO) -> Iterator[dict]:
    """Yield raw rows from the documented workbook; close file resources reliably."""
    from openpyxl import load_workbook
    if hasattr(source, "seek"):
        source.seek(0)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if EXCEL_SHEET not in workbook.sheetnames:
            raise ValueError(f"Feuille requise absente : {EXCEL_SHEET}.")
        rows = iter(workbook[EXCEL_SHEET].iter_rows(values_only=True))
        headers = list(next(rows, []))
        if headers != EXCEL_HEADERS:
            raise ValueError("Colonnes attendues dans l’ordre : User, Q1, Q2, …, Q134.")
        seen = set()
        for row in rows:
            if not row or all(value is None for value in row):
                continue
            values = list(row) + [None] * max(0, len(headers) - len(row))
            record = dict(zip(headers, values))
            user = record["User"]
            if _missing(user):
                raise ValueError("Une ligne contient des réponses sans identifiant User.")
            record["User"] = str(user)
            if record["User"] in seen:
                raise ValueError(f"Identifiant User dupliqué : {record['User']}.")
            seen.add(record["User"])
            yield record
    finally:
        workbook.close()


def list_excel_users(source: Path | BytesIO, limit: int = 10000) -> list[str]:
    if not isinstance(limit, int) or isinstance(limit, bool) or limit < 0:
        raise ValueError("La limite doit être un entier positif ou nul.")
    if limit == 0:
        return []
    rows = iter_excel_profiles(source)
    try:
        users = []
        for record in rows:
            users.append(record["User"])
            if len(users) >= limit:
                break
        return users
    finally:
        rows.close()


def load_excel_dataset(source: Path | BytesIO) -> dict[str, dict]:
    """Read the uploaded answer workbook once for fast User selection.

    Workbook order and exact User identifiers are preserved. Values stay raw
    until the selected row is passed to build_profile(). The mapping is meant
    for the current session only: nothing is serialized or written to disk.
    Reading the complete iterator validates headers, missing IDs and duplicate
    IDs even when the first user would otherwise already be selectable.
    """
    rows = iter_excel_profiles(source)
    try:
        profiles = {record["User"]: record for record in rows}
    finally:
        rows.close()
    if not profiles:
        raise ValueError("Le classeur ne contient aucun profil User à sélectionner.")
    return profiles


def load_excel_profile(source: Path | BytesIO, user_id: str) -> dict:
    rows = iter_excel_profiles(source)
    try:
        for record in rows:
            if record["User"] == str(user_id):
                return record
    finally:
        rows.close()
    raise KeyError(f"Profil introuvable : {user_id}")
